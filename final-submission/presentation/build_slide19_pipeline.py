"""Label the PaperBanana artwork and optionally place it on slide 19.

The artwork is illustrative. All process/status labels below are drawn from
PPT_BUILD_GUIDE.md and DEMO_RUNBOOK.md, not inferred from the generated image.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont


HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
BASE = ASSETS / "slide19-pipeline-paperbanana-base.png"
FINAL = ASSETS / "slide19-synthetic-matter-pipeline.png"
DECK = HERE / "Draftly-Final-Academic-Presentation-Reviewed.pptx"
OUT_DECK = HERE / "Draftly-Final-Academic-Presentation-Slide19-Pipeline.pptx"
PDF = HERE / "Draftly-Final-Academic-Presentation-Reviewed.pdf"
OUT_PDF = HERE / "Draftly-Final-Academic-Presentation-Slide19-Pipeline.pdf"

BLUE = "#0737B5"
NAVY = "#17335E"
PALE = "#E9F0FF"
INK = "#172332"
GREY = "#5D6978"
AMBER = "#D99300"


def font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / name), size)


def centered(draw: ImageDraw.ImageDraw, x: int, y: int, value: str,
             face: ImageFont.FreeTypeFont, fill: str) -> None:
    box = draw.textbbox((0, 0), value, font=face)
    draw.text((x - (box[2] - box[0]) / 2, y), value, font=face, fill=fill)


def build() -> None:
    art = Image.open(BASE).convert("RGB")
    # The official renderer put duplicated text above and below the icons.
    # Keep only its illustrated row; place verified labels in a clean layer.
    icons = art.crop((0, 600, 2752, 950)).resize((2860, 364), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (3000, 1160), "#FFFFFF")
    draw = ImageDraw.Draw(canvas)

    groups = [
        (80, 1665, "EVIDENCE PATH"),
        (1694, 2077, "MATTER CHAT"),
        (2092, 2470, "RESEARCH VIEW"),
        (2485, 2920, "PREFLIGHT GATE"),
    ]
    for left, right, title in groups:
        draw.rounded_rectangle((left, 54, right, 146), radius=24, fill=PALE)
        centered(draw, (left + right) // 2, 77, title, font(35, bold=True), BLUE)

    canvas.paste(icons, (70, 240))
    labels = [
        (342, "Matter", "persistent record"),
        (733, "Documents", "source file"),
        (1103, "Candidate review", "lawyer verifies"),
        (1487, "Checks", "rule-pack finding"),
        (1899, "Matter chat", "matter state only"),
        (2262, "Cited research", "open separately"),
        (2646, "Draft / preflight", "refusal if blocked"),
    ]
    for x, heading, detail in labels:
        centered(draw, x, 651, heading, font(45, bold=True), INK)
        centered(draw, x, 718, detail, font(35), GREY)

    draw.rounded_rectangle((79, 849, 1761, 1065), radius=31, fill=PALE)
    draw.text((134, 885), "HOSTED DEMO STATUS", font=font(35, bold=True), fill=BLUE)
    draw.text((134, 951), "Synthetic matter  ·  stub extraction  ·  BM25 search",
              font=font(45), fill=NAVY)

    draw.rounded_rectangle((1792, 849, 2920, 1065), radius=31,
                           fill="#FFF3DC", outline="#DCAA49", width=3)
    draw.text((1845, 885), "DECISION GATE", font=font(35, bold=True), fill="#925D00")
    draw.text((1845, 951), "Unresolved blocker  →  stop", font=font(45, bold=True), fill="#744B00")

    canvas.save(FINAL, optimize=True)
    print(f"Saved {FINAL} ({canvas.width} x {canvas.height})")


def place() -> None:
    import fitz
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation(DECK)
    if len(presentation.slides) != 22:
        raise RuntimeError(f"Expected 22 slides, found {len(presentation.slides)}")
    slide = presentation.slides[18]
    # Overlay the old route card only. Keep the existing slide title, logo,
    # page number, and speaker notes intact.
    slide.shapes.add_picture(str(FINAL), Inches(0.63), Inches(1.84),
                             width=Inches(12.06), height=Inches(4.66))
    presentation.save(OUT_DECK)
    print(f"Saved 22-slide deck with new slide 19: {OUT_DECK}")

    # The reviewed PDF is the already-rendered source deck. Applying the same
    # image rectangle here gives a distributable preview without invoking the
    # PowerPoint GUI or changing other slides.
    pdf = fitz.open(PDF)
    if len(pdf) != 22:
        raise RuntimeError(f"Expected 22 PDF pages, found {len(pdf)}")
    page = pdf[18]
    box = fitz.Rect(0.63 * 72, 1.84 * 72, (0.63 + 12.06) * 72,
                    (1.84 + 4.66) * 72)
    page.insert_image(box, filename=str(FINAL), overlay=True)
    pdf.save(OUT_PDF, garbage=4, deflate=True)
    pdf.close()
    print(f"Saved matching 22-page PDF preview: {OUT_PDF}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--place", action="store_true", help="Overlay slide 19 in the reviewed PPTX")
    args = parser.parse_args()
    build()
    if args.place:
        place()
