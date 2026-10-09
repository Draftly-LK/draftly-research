"""Build the editable Draftly final presentation from PPT_BUILD_GUIDE.md.

The visual system follows Guizang Style B (Swiss, International Klein Blue).
All audience slides are native PowerPoint shapes/text plus embedded source figures.
Speaker notes are populated from the single slide guide.
"""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
FIG = HERE.parent / "figures"
ASSET = HERE / "assets"
GUIDE = HERE / "PPT_BUILD_GUIDE.md"
OUTPUT = HERE / "Draftly-Final-Academic-Presentation-Reviewed.pptx"

W, H = 13.333, 7.5
PAPER = "FAFAF8"
INK = "0A0A0A"
GREY1 = "F0F0EE"
GREY2 = "D4D4D2"
GREY3 = "737373"
BLUE = "002FA7"
BLUE_PALE = "E8EEFC"
WHITE = "FFFFFF"
RED = "A3292A"
FONT = "Arial"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def clean(value: str) -> str:
    value = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value)
    return value.replace("**", "").replace("*", "").replace("  ", " ").strip()


def parse_guide() -> list[dict]:
    source = GUIDE.read_text(encoding="utf-8")
    chunks = re.split(r"(?=^## Slide \d+ — )", source, flags=re.M)[1:]
    slides = []
    for chunk in chunks:
        head = re.match(r"## Slide (\d+) — `([^`]+)` — (.+)", chunk)
        if not head:
            continue
        copy_block = re.search(r"### On-slide copy\s*\n(.*?)(?=\n\*\*Speaker notes:\*\*)", chunk, re.S)
        notes = re.search(r"\*\*Speaker notes:\*\* [“\"](.*?)[”\"]\s*\n", chunk, re.S)
        time = re.search(r"\*\*Time:\*\* (\d+:\d+)", chunk)
        if not copy_block or not notes or not time:
            raise ValueError(f"Missing required guide field on slide {head.group(1)}")
        copy = [clean(line[1:]) for line in copy_block.group(1).splitlines() if line.startswith(">")]
        slides.append({
            "number": int(head.group(1)),
            "id": head.group(2),
            "title": head.group(3).strip(),
            "copy": copy,
            "notes": notes.group(1).strip(),
            "time": time.group(1),
        })
    assert len(slides) == 22, len(slides)
    return slides


def rect(slide, x, y, w, h, fill=PAPER, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def line(slide, x1, y1, x2, y2, color=GREY2, width=1.0):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def txt(slide, value, x, y, w, h, size=18, color=INK, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, para_text in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = para_text
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box


def image_fit(slide, path: Path, x, y, w, h, border=False):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w-pw)/2, y + (h-ph)/2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(pw), height=Inches(ph))
    if border:
        pic.line.color.rgb = rgb(GREY2)
        pic.line.width = Pt(0.7)
    return pic


def image_crop(slide, path: Path, x, y, w, h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    iw, ih = Image.open(path).size
    image_aspect = iw/ih
    box_aspect = w/h
    if image_aspect > box_aspect:
        frac = 1 - box_aspect/image_aspect
        pic.crop_left = pic.crop_right = frac/2
    else:
        frac = 1 - image_aspect/box_aspect
        pic.crop_top = pic.crop_bottom = frac/2
    return pic


def logo(slide, dark=False, large=False):
    path = ASSET / ("draftly-logo-white-transparent.png" if dark else "draftly-logo-blue-transparent.png")
    if large:
        return image_fit(slide, path, 0.53, 0.48, 2.2, 1.35)
    return image_fit(slide, path, 11.98, 0.22, 0.80, 0.53)


def base(prs: Presentation, item: dict, section="", dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, fill=BLUE if dark else PAPER)
    if item["number"] not in (1, 22):
        txt(slide, section.upper() or "DRAFTLY · FINAL PRESENTATION", 0.55, 0.29, 10.8, 0.23,
            size=9.5, color=BLUE if not dark else WHITE, bold=True)
        txt(slide, item["title"], 0.55, 0.63, 11.25, 0.88, size=29,
            color=INK if not dark else WHITE, valign=MSO_ANCHOR.TOP)
        logo(slide, dark=dark)
        line(slide, 0.55, 7.10, 12.78, 7.10, color=GREY2 if not dark else WHITE, width=0.6)
        txt(slide, "DRAFTLY  /  UNIVERSITY OF MORATUWA", 0.55, 7.17, 6.0, 0.17,
            size=8.3, color=GREY3 if not dark else WHITE)
        txt(slide, f"{item['number']:02d} / 22", 11.70, 7.16, 1.08, 0.18,
            size=8.3, color=GREY3 if not dark else WHITE, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = (
        f"Slide {item['number']:02d} · {item['id']} · {item['time']}\n\n"
        + item["notes"]
    )
    return slide


def label(slide, text, x, y, w, color=GREY3):
    return txt(slide, text.upper(), x, y, w, 0.25, size=10, color=color, bold=True)


def number_card(slide, number, caption, x, y, w, h=1.25, color=BLUE):
    rect(slide, x, y, w, h, fill=WHITE, line=GREY2)
    txt(slide, number, x+0.20, y+0.12, w-0.40, 0.64, size=31, color=color)
    txt(slide, caption, x+0.20, y+0.78, w-0.40, h-0.87, size=12.3, color=INK, valign=MSO_ANCHOR.TOP)


def chip(slide, value, x, y, w, h=0.38, fill=BLUE_PALE, color=BLUE, size=12):
    rect(slide, x, y, w, h, fill=fill)
    txt(slide, value, x+0.1, y+0.03, w-0.2, h-0.06, size=size, color=color, bold=True)


def flow(slide, labels, x, y, w, h, emphasis=None, font_size=14):
    gap = 0.23
    each = (w - (len(labels)-1)*gap) / len(labels)
    for idx, text in enumerate(labels):
        bx = x + idx*(each+gap)
        active = emphasis is not None and idx in emphasis
        rect(slide, bx, y, each, h, fill=BLUE if active else WHITE, line=None if active else GREY2)
        txt(slide, text, bx+0.10, y+0.10, each-0.20, h-0.20, size=font_size,
            color=WHITE if active else INK, bold=active, align=PP_ALIGN.CENTER)
        if idx < len(labels)-1:
            txt(slide, "→", bx+each, y+h/2-0.15, gap, 0.3, size=15, color=BLUE, align=PP_ALIGN.CENTER)


def build():
    items = parse_guide()
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Draftly · Final Academic Presentation"
    prs.core_properties.subject = "Sri Lankan legal retrieval and a lawyer-controlled conveyancing workbench"
    prs.core_properties.author = "Draftly Project Team"

    for item in items:
        n = item["number"]
        section = "Part I · Legal information retrieval" if 4 <= n <= 11 else (
            "Part II · Software engineering" if 12 <= n <= 20 else "Draftly · Final presentation")
        dark = n in (1, 22)
        slide = base(prs, item, section, dark=dark)

        if n == 1:
            logo(slide, dark=True, large=True)
            txt(slide, "DRAFTLY", 0.72, 2.15, 11.65, 1.0, size=59, color=WHITE)
            txt(slide, "Grounded legal retrieval and a lawyer-controlled\nconveyancing matter workflow for Sri Lanka",
                0.72, 3.25, 11.25, 1.18, size=27, color=WHITE)
            line(slide, 0.72, 5.16, 12.3, 5.16, color=WHITE, width=0.7)
            txt(slide, "GROUP 06  ·  UNIVERSITY OF MORATUWA  ·  DCS&E", 0.72, 5.36, 11.25, 0.35, size=14, color=WHITE)
            txt(slide, "Himath Dhanapala  ·  Lahiru Dilshan  ·  Praveen De Silva", 0.72, 5.87, 11.25, 0.36, size=17, color=WHITE)
            txt(slide, "Supervisor: Dr. Nisansa de Silva", 0.72, 6.35, 10, 0.28, size=13, color=WHITE)

        elif n == 2:
            for x, num, heading, desc in [
                (0.62, "01", "LEGAL INFORMATION RETRIEVAL", "One scenario can require several provisions. A missing section leaves the answer incomplete."),
                (6.81, "02", "MATTER DOCUMENTS", "Key facts are buried in PDFs and scans. Lawyers must read and reconcile them."),
            ]:
                rect(slide, x, 1.86, 5.89, 3.75, fill=WHITE, line=GREY2)
                txt(slide, num, x+0.28, 2.12, 1.4, 0.9, size=52, color=BLUE)
                txt(slide, heading, x+0.28, 3.12, 5.25, 0.48, size=18, color=INK, bold=True)
                txt(slide, desc, x+0.28, 3.75, 5.25, 1.22, size=21, color=INK, valign=MSO_ANCHOR.TOP)
            chip(slide, "MEASURED LEGAL SEARCH  +  SOURCE-LINKED DOCUMENT REVIEW", 0.62, 6.01, 12.08, 0.54, fill=BLUE, color=WHITE, size=16)

        elif n == 3:
            portrait = ASSET / "consultations" / "anura-dhanaratna.jpeg"
            image_crop(slide, portrait, 0.62, 1.80, 2.35, 3.75)
            txt(slide, "ANURA DHANARATNA", 0.62, 5.65, 2.5, 0.30, size=13, color=BLUE, bold=True)
            txt(slide, "Lawyer · Notary · Law College lecturer", 0.62, 6.00, 3.10, 0.37, size=11.5)
            chip(slide, "3 IN-PERSON CONSULTATIONS", 3.35, 1.83, 4.55, 0.43, fill=BLUE, color=WHITE, size=14)
            rows = [
                ("Notarial curriculum", "Topic map and statutory coverage"),
                ("Sample matter documents", "Document-reading trials"),
                ("Practitioner feedback", "Evidence review and lawyer decision gates"),
            ]
            for idx, (a, b) in enumerate(rows):
                y = 2.50+idx*1.05
                rect(slide, 3.35, y, 9.28, 0.90, fill=WHITE, line=GREY2)
                txt(slide, a, 3.58, y+0.12, 3.45, 0.64, size=17, color=BLUE, bold=True)
                txt(slide, "→", 7.00, y+0.16, 0.34, 0.5, size=21, color=BLUE)
                txt(slide, b, 7.35, y+0.12, 4.97, 0.64, size=17)
            label(slide, "Additional stakeholder conversations", 8.31, 5.90, 3.0)
            for idx, name in enumerate(["ishan-rathnapala.jpeg", "priyal-wijayaweera.jpeg", "stakeholder-portrait-04.jpeg"]):
                image_crop(slide, ASSET/"consultations"/name, 11.48-idx*0.72, 5.87, 0.56, 0.65)

        elif n == 4:
            tight = ASSET/"research-statutory-corpus-tight.png"
            Image.open(ASSET/"research-statutory-corpus-crop.png").crop((0, 0, 1490, 1168)).save(tight)
            image_fit(slide, tight, 0.62, 1.78, 5.48, 4.80)
            label(slide, "Internal parsed research corpus", 0.67, 6.56, 5.2)
            number_card(slide, "52 + 60", "principal enactments + amending Acts", 6.44, 1.88, 5.90, 1.30)
            number_card(slide, "4,557", "sections · 17,854 provisions", 6.44, 3.30, 5.90, 1.30)
            number_card(slide, "6,346", "typed statutory edges", 6.44, 4.72, 5.90, 1.30)
            hf_mark = image_fit(slide, ASSET/"hugging-face-logo.png", 6.45, 6.25, 0.38, 0.36)
            hf_mark.click_action.hyperlink.address = "https://huggingface.co/datasets/lanka-legal-nlp/draftly-sri-lanka-act-sources"
            txt(slide, "Public Act source-link index · parsed text stays internal", 6.89, 6.27, 5.7, 0.28, size=12, color=GREY3)

        elif n == 5:
            number_card(slide, "3,703", "reported NLR / SLR conveyancing judgments", 0.62, 1.85, 5.88, 2.32)
            number_card(slide, "5,474", "official Supreme Court / Court of Appeal judgment documents", 6.83, 1.85, 5.88, 2.32)
            line(slide, 0.62, 4.48, 12.71, 4.48, color=GREY2)
            txt(slide, "3,521 / 3,703", 0.64, 4.73, 5.8, 0.80, size=43, color=BLUE)
            txt(slide, "reported cases yielded verbatim headnote rules", 0.64, 5.59, 8.9, 0.52, size=21)
            chip(slide, "PROVISIONAL RULES  ·  LAWYER VALIDATION PENDING", 0.64, 6.37, 7.8, 0.41)

        elif n == 6:
            chip(slide, "5,121 CONVEYANCING-FLAGGED CASES  /  9,177 COLLECTED", 0.62, 1.77, 9.25, 0.44, fill=BLUE, color=WHITE, size=16)
            rect(slide, 0.62, 2.73, 2.00, 1.40, fill=BLUE)
            txt(slide, "Fact pattern", 0.81, 3.18, 1.64, 0.45, size=20, color=WHITE, align=PP_ALIGN.CENTER)
            txt(slide, "→", 2.63, 3.18, 0.35, 0.35, size=19, color=BLUE, align=PP_ALIGN.CENTER)
            rect(slide, 3.00, 2.73, 2.02, 1.40, fill=WHITE, line=GREY2)
            txt(slide, "Case index", 3.18, 3.18, 1.68, 0.45, size=19, align=PP_ALIGN.CENTER)
            line(slide, 5.03, 3.43, 5.38, 3.43, color=BLUE, width=1.4)
            line(slide, 5.38, 2.49, 5.38, 4.40, color=BLUE, width=1.4)
            branches = [
                ("BM25 judgment / rule text", 2.31, WHITE, INK),
                ("Case graph: statutes / topics / catchwords", 3.23, BLUE_PALE, BLUE),
                ("Dense embeddings: optional, inactive in pilot", 4.15, GREY1, GREY3),
            ]
            for value, yy, fill, color in branches:
                line(slide, 5.38, yy+0.33, 5.58, yy+0.33, color=BLUE if color != GREY3 else GREY2, width=1.2)
                rect(slide, 5.58, yy, 3.10, 0.68, fill=fill, line=GREY2 if fill==WHITE else None)
                txt(slide, value, 5.75, yy+0.10, 2.77, 0.47, size=14, color=color)
                line(slide, 8.68, yy+0.33, 8.89, yy+0.33, color=BLUE if color != GREY3 else GREY2, width=1.2)
            line(slide, 8.89, 2.64, 8.89, 4.48, color=BLUE, width=1.4)
            txt(slide, "→", 8.91, 3.19, 0.31, 0.40, size=19, color=BLUE, align=PP_ALIGN.CENTER)
            rect(slide, 9.24, 2.73, 1.57, 1.40, fill=BLUE_PALE)
            txt(slide, "Rank fusion + corroboration", 9.37, 3.01, 1.31, 0.85, size=14.2, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
            txt(slide, "→", 10.82, 3.19, 0.31, 0.40, size=19, color=BLUE, align=PP_ALIGN.CENTER)
            rect(slide, 11.16, 2.73, 1.55, 1.40, fill=BLUE)
            txt(slide, "Cited cases or none", 11.28, 3.06, 1.31, 0.71, size=14.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            rect(slide, 0.62, 5.28, 12.08, 1.13, fill=WHITE, line=GREY2)
            txt(slide, "20-query proxy pilot", 0.84, 5.53, 3.49, 0.30, size=14, color=BLUE, bold=True)
            txt(slide, "14 / 20 on the initial corpus snapshot", 4.04, 5.46, 3.45, 0.48, size=19, color=INK)
            txt(slide, "10 / 20 after rebaseline", 8.30, 5.46, 3.91, 0.48, size=19, color=INK)
            txt(slide, "LLM-judged relevance · no attorney labels", 0.84, 6.51, 8.38, 0.23, size=11.2, color=GREY3)

        elif n == 7:
            for idx, (top, bottom) in enumerate([
                ("01 · QUESTION SELECTION", "16 papers → 667 atomic questions"),
                ("02 · PROVISION LABELS", "50 scored questions / 40 matters"),
                ("03 · HUMAN REVIEW", "Three-attorney stage pending"),
            ]):
                x = 0.62+idx*4.07
                rect(slide, x, 1.80, 3.82, 0.78, fill=BLUE if idx<2 else GREY1)
                txt(slide, top, x+0.14, 1.90, 3.54, 0.22, size=10.6, color=WHITE if idx<2 else GREY3, bold=True)
                txt(slide, bottom, x+0.14, 2.15, 3.54, 0.31, size=14.3, color=WHITE if idx<2 else INK)
            image_fit(slide, FIG/"research-benchmark-construction-preview.png", 0.62, 2.74, 12.08, 3.70)
            chip(slide, "40 TEST QUESTIONS / 32 MATTERS  ·  PROVISIONAL LABELS", 0.62, 6.47, 9.70, 0.43, fill=BLUE_PALE, color=BLUE, size=13)
            hf_mark = image_fit(slide, ASSET/"hugging-face-logo.png", 10.53, 6.48, 0.40, 0.40)
            hf_mark.click_action.hyperlink.address = "https://huggingface.co/datasets/lanka-legal-nlp/draftly-statutory-retrieval/tree/v1.0-paper-provisional"
            txt(slide, "Provisional release", 10.99, 6.52, 1.75, 0.27, size=11.5, color=BLUE, bold=True)

        elif n == 8:
            label(slide, "Illustrative question: agent-proposed reference bundle", 0.64, 1.77, 9.6)
            for idx, lab in enumerate(["Section A", "Section B", "Section C"]):
                x = 0.64+idx*2.05
                rect(slide, x, 2.27, 1.75, 1.12, fill=BLUE if idx < 2 else WHITE, line=None if idx < 2 else BLUE)
                txt(slide, lab, x+0.10, 2.62, 1.55, 0.30, size=18, color=WHITE if idx < 2 else BLUE, align=PP_ALIGN.CENTER)
            txt(slide, "→", 6.78, 2.58, 0.9, 0.44, size=35, color=BLUE, align=PP_ALIGN.CENTER)
            rect(slide, 8.08, 2.14, 4.24, 1.54, fill=GREY1)
            txt(slide, "TOP 20", 8.36, 2.38, 3.65, 0.25, size=11, color=GREY3, bold=True)
            txt(slide, "A  +  B", 8.36, 2.76, 3.65, 0.47, size=29, color=INK)
            chip(slide, "SECTION C MISSING", 8.08, 4.20, 4.24, 0.49, fill="F7E9E9", color=RED, size=17)
            txt(slide, "C@20 = 0", 0.64, 4.37, 6.6, 1.02, size=52, color=BLUE)
            txt(slide, "Full credit only when every proposed section appears in the first 20 results.",
                0.64, 5.54, 11.71, 0.74, size=23)

        elif n == 9:
            methods_crop = ASSET/"research-retrieval-methods-top.png"
            Image.open(ASSET/"research-retrieval-system-crop.png").crop((0, 0, 2075, 965)).save(methods_crop)
            image_fit(slide, methods_crop, 0.62, 1.79, 8.02, 4.68)
            for idx, (a, b) in enumerate([
                ("LEXICAL", "BM25 · field-weighted BM25"),
                ("SEMANTIC / FUSION", "Dense · hybrid · reranker"),
                ("STRUCTURE", "Hierarchical · typed expansion"),
            ]):
                y = 1.87+idx*1.34
                rect(slide, 8.85, y, 3.85, 1.14, fill=WHITE, line=GREY2)
                label(slide, a, 9.06, y+0.14, 3.4, color=BLUE)
                txt(slide, b, 9.06, y+0.48, 3.45, 0.44, size=15)
            chip(slide, "OFFLINE RESEARCH · PROVISIONAL LABELS", 8.85, 6.03, 3.85, 0.42, fill=BLUE, color=WHITE, size=11.5)
            txt(slide, "Seven matched configurations · one frozen test split", 0.64, 6.65, 10.9, 0.25, size=12, color=GREY3)

        elif n == 10:
            image_fit(slide, HERE/"retrieval-c20-results-slide.png", 0.55, 1.80, 8.15, 4.84)
            number_card(slide, "0 / 20", "questions needing ≥4 provisions completed", 8.91, 1.86, 3.75, 1.53)
            number_card(slide, "2 / 114", "missed provisions one explicit edge away", 8.91, 3.59, 3.75, 1.53)
            rect(slide, 8.91, 5.36, 3.75, 1.07, fill=BLUE_PALE)
            txt(slide, "No clear winner among the leading systems", 9.10, 5.53, 3.38, 0.73, size=16, color=BLUE)
            txt(slide, "OFFLINE · PROVISIONAL LABELS · 40 TEST QUESTIONS / 32 MATTERS", 0.62, 6.70, 11.9, 0.24, size=10.5, color=GREY3, bold=True)

        elif n == 11:
            image_fit(slide, ASSET/"accepted-paper-title-authors.png", 0.67, 1.91, 7.07, 3.08)
            chip(slide, "ACCEPTED FOR PRESENTATION", 0.67, 5.25, 7.07, 0.46, fill=BLUE, color=WHITE, size=17)
            txt(slide, "NeurIPS 2026 GlobalSouthAI workshop", 0.67, 5.83, 7.08, 0.42, size=16, color=BLUE)
            for idx, value in enumerate(["Structured Sri Lankan statutory corpus", "Provisional scenario QA benchmark", "Seven-system retrieval evaluation"]):
                rect(slide, 8.10, 1.91+idx*1.16, 4.47, 0.95, fill=WHITE, line=GREY2)
                txt(slide, f"0{idx+1}", 8.34, 2.10+idx*1.16, 0.49, 0.40, size=23, color=BLUE)
                txt(slide, value, 8.89, 2.08+idx*1.16, 3.40, 0.50, size=16)
            txt(slide, "Typed-edge expansion did not demonstrate a complete-recall gain over hybrid retrieval.",
                8.10, 5.50, 4.47, 0.74, size=16, color=INK)

        elif n == 12:
            chip(slide, "FIRST SCOPE · REGISTRATION OF TITLE ACT (RTA)", 0.63, 1.80, 8.78, 0.49, fill=BLUE, color=WHITE, size=16)
            rect(slide, 0.63, 2.52, 8.77, 3.67, fill=BLUE_PALE)
            flow(slide, ["Title certificate", "Cadastral / survey plan", "Transaction evidence"],
                 0.95, 2.99, 8.12, 1.16, emphasis={0, 2}, font_size=15)
            txt(slide, "Defined requirements + prescribed forms", 0.96, 4.70, 8.08, 0.55, size=23, color=BLUE)
            txt(slide, "support reviewable checks", 0.96, 5.35, 7.92, 0.46, size=17, color=BLUE)
            label(slide, "Deferred", 9.70, 1.87, 2.9)
            txt(slide, "Handwritten pattiru / folio OCR\n\nAT-form extraction\n\nHistorical title reconstruction\n\nInitial title settlement",
                9.70, 2.53, 2.63, 3.76, size=16, color=GREY3, valign=MSO_ANCHOR.TOP)

        elif n == 13:
            lanes = [
                ("01 · MATTER AND EVIDENCE", ["Create + classify matter", "Upload source documents", "Record candidate facts + evidence"]),
                ("02 · LAWYER REVIEW", ["Verify or correct facts", "Run checklist + checks", "Resolve finding / request evidence", "Matter Agent conversation", "Cited research (separate view)"]),
                ("03 · DRAFT AND RECORD", ["Prepare form draft", "Check prerequisites", "Lawyer reviews + decides", "Record approval / export + audit"]),
            ]
            for idx, (heading, steps) in enumerate(lanes):
                x = 0.62 + idx*4.10
                rect(slide, x, 1.92, 3.83, 4.59, fill=WHITE, line=GREY2)
                rect(slide, x, 1.92, 3.83, 0.62, fill=BLUE)
                txt(slide, heading, x+0.18, 2.09, 3.49, 0.31, size=13.5, color=WHITE, bold=True)
                height = 0.61 if len(steps) == 5 else (0.72 if len(steps) == 4 else 0.85)
                gap = 0.14 if len(steps) == 5 else 0.23
                start_y = 2.73 if len(steps) == 5 else 2.84
                for j, step in enumerate(steps):
                    yy = start_y + j*(height+gap)
                    fill = BLUE_PALE if (idx == 1 and j in (0, 2)) or (idx == 2 and j == 2) else GREY1
                    rect(slide, x+0.19, yy, 3.45, height, fill=fill)
                    txt(slide, step, x+0.35, yy+0.10, 3.12, height-0.18, size=14.2,
                        color=BLUE if fill == BLUE_PALE else INK)
                    if j < len(steps)-1:
                        txt(slide, "↓", x+1.77, yy+height+0.01, 0.30, gap-0.03, size=13, color=BLUE, align=PP_ALIGN.CENTER)
                if idx < 2:
                    txt(slide, "→", x+3.82, 3.99, 0.28, 0.35, size=21, color=BLUE, align=PP_ALIGN.CENTER)
            txt(slide, "Findings return to review; a rejected draft returns to verified facts.",
                0.63, 6.66, 11.55, 0.26, size=12.0, color=GREY3)

        elif n == 14:
            image_fit(slide, ASSET/"redacted-real-ocr-boxed-page.png", 0.66, 1.78, 4.04, 4.85)
            txt(slide, "Real benchmark page · text redacted · OCR regions shown", 0.66, 6.67, 5.1, 0.24, size=10.5, color=GREY3)
            number_card(slide, "4 matters · 38 documents", "282 pages in the benchmark inventory", 5.12, 1.87, 7.30, 1.15)
            number_card(slide, "26 / 26 pages", "pilot pages returned text", 5.12, 3.14, 7.30, 1.15)
            number_card(slide, "37 labelled fields", "across 4 documents in 1 matter", 5.12, 4.41, 7.30, 1.15)
            chip(slide, "FIELD-LEVEL ACCURACY NOT YET SCORED", 5.12, 5.85, 7.30, 0.57, fill=BLUE, color=WHITE, size=17)

        elif n == 15:
            stages_crop = ASSET/"document-processing-primary-stages.png"
            Image.open(ASSET/"document-processing-research-pipeline-crop.png").crop((0, 0, 2752, 890)).save(stages_crop)
            for idx, (heading, detail) in enumerate([
                ("01 · CAPTURE", "Upload · OCR · quality check"),
                ("02 · ORGANIZE", "Rotate · classify · group"),
                ("03 · PROPOSE + VERIFY", "Derivatives · fields · human review"),
            ]):
                x = 0.62+idx*4.08
                rect(slide, x, 1.77, 3.83, 0.73, fill=BLUE if idx == 0 else BLUE_PALE)
                txt(slide, heading, x+0.14, 1.87, 3.54, 0.22, size=10.5, color=WHITE if idx == 0 else BLUE, bold=True)
                txt(slide, detail, x+0.14, 2.13, 3.54, 0.25, size=13.8, color=WHITE if idx == 0 else INK)
            image_fit(slide, stages_crop, 0.62, 2.55, 12.08, 3.88)
            chip(slide, "RESEARCH DESIGN  ·  LATER EXTRACTION STAGES ARE NOT HOSTED", 0.62, 6.45, 12.08, 0.33, fill=BLUE_PALE, color=BLUE, size=12.6)
            txt(slide, "Hosted today: upload → processing record / stub candidates → lawyer review", 0.62, 6.84, 11.72, 0.20, size=11.3, color=GREY3)

        elif n == 16:
            chip(slide, "HOSTED RESEARCH PATH · BM25", 0.62, 1.80, 6.17, 0.46, fill=BLUE, color=WHITE, size=17)
            flow(slide, ["Lawyer question", "Research view", "FastAPI", "BM25 index", "Cited passages"],
                 0.62, 2.78, 12.03, 1.38, emphasis={0, 3}, font_size=15)
            rect(slide, 0.62, 4.69, 7.41, 1.33, fill=BLUE_PALE)
            txt(slide, "Grounded response\nor insufficient authority", 0.86, 4.86, 6.92, 0.85, size=22, color=BLUE)
            rect(slide, 8.34, 4.69, 4.31, 1.33, fill=WHITE, line=GREY2)
            label(slide, "Current boundary", 8.59, 4.87, 3.78, color=BLUE)
            txt(slide, "Separate research workspace; no direct Matter Agent search", 8.59, 5.21, 3.72, 0.59, size=15.3)
            txt(slide, "Next version in progress: embeddings + query rewriting", 0.62, 6.56, 10.80, 0.27, size=12.5, color=GREY3)

        elif n == 17:
            cropped = ASSET / "matter-agent-architecture-slide-crop.png"
            im = Image.open(FIG/"platform-matter-agent-service-architecture.png")
            im.crop((610, 0, 2752, 1090)).save(cropped)
            image_fit(slide, cropped, 0.61, 1.84, 8.10, 4.81)
            label(slide, "Earlier design reference", 0.63, 6.64, 7.55)
            for idx, value in enumerate([
                "Matter-scoped session",
                "Allowlisted server tools",
                "Action → explicit confirmation",
                "Legal question → abstention",
                "Cited research → separate view",
            ]):
                rect(slide, 8.93, 1.84+idx*0.94, 3.73, 0.76, fill=WHITE, line=GREY2)
                txt(slide, value, 9.10, 1.99+idx*0.94, 3.38, 0.48, size=14.4,
                    color=BLUE if idx in (2, 3) else INK, bold=idx in (2, 3))

        elif n == 18:
            flow(slide, ["Versioned RTA\nrule pack", "Checklist +\nfindings", "Verified facts +\nform bindings", "Preflight +\nlawyer approval"],
                 0.64, 2.08, 12.04, 1.67, emphasis={0, 3}, font_size=16)
            rect(slide, 3.60, 4.42, 6.14, 1.27, fill="F7E9E9")
            txt(slide, "UNRESOLVED STATUTORY BLOCKER  →  STOP", 3.86, 4.79, 5.61, 0.57,
                size=20, color=RED, bold=True, align=PP_ALIGN.CENTER)
            txt(slide, "An unapproved draft cannot be exported as approved.", 2.26, 6.12, 8.82, 0.39,
                size=18, color=INK, align=PP_ALIGN.CENTER)

        elif n == 19:
            chip(slide, "LIVE PRODUCT DEMONSTRATION  ·  SYNTHETIC RTA MATTER", 0.63, 1.84, 9.93, 0.48, fill=BLUE, color=WHITE, size=17)
            flow(slide, ["Matter", "Documents", "Candidate review", "Checks", "Matter chat", "Cited research", "Draft / preflight"],
                 0.63, 2.92, 12.06, 1.47, emphasis={0, 2, 6}, font_size=12.4)
            rect(slide, 0.63, 4.99, 8.04, 1.38, fill=BLUE_PALE)
            txt(slide, "Show the refusal gate and the evidence trail.", 0.90, 5.29, 7.52, 0.66, size=24, color=BLUE)
            rect(slide, 8.96, 4.99, 3.73, 1.38, fill=WHITE, line=GREY2)
            txt(slide, "HOSTED STATUS", 9.17, 5.16, 3.28, 0.22, size=10, color=BLUE, bold=True)
            txt(slide, "Stub extraction · BM25 search", 9.17, 5.54, 3.25, 0.49, size=16)
            txt(slide, "Switch to the rehearsed live app or verified recording.", 0.63, 6.64, 10.4, 0.24, size=11, color=GREY3)

        elif n == 20:
            table = [
                ("STATUTORY IR", "C@20 0.375 · 40 test questions", "Provisional labels"),
                ("SIMILAR CASES", "14/20 initial · 10/20 rebaseline", "Proxy judged; dense inactive"),
                ("CASE-LAW RULES", "74/100 quote-grounded", "58.1% proxy usable"),
                ("DOCUMENT READING", "26/26 pilot pages returned text", "Field accuracy unscored"),
                ("PLATFORM", "3,906 backend + 348 frontend passed", "3 browser refusals; one journey"),
            ]
            for idx, (name, measure, limit) in enumerate(table):
                y = 1.82+idx*0.88
                rect(slide, 0.62, y, 12.05, 0.73, fill=WHITE if idx%2 == 0 else GREY1)
                txt(slide, name, 0.83, y+0.13, 2.10, 0.44, size=12.5, color=BLUE, bold=True)
                txt(slide, measure, 3.02, y+0.13, 5.49, 0.44, size=15.2)
                txt(slide, limit, 8.65, y+0.13, 3.82, 0.44, size=14.4, color=INK)
            chip(slide, "WIDER BROWSER RUN: 4 PASSED  ·  9 FAILED  ·  2 SKIPPED", 0.62, 6.48, 12.05, 0.43,
                 fill=BLUE_PALE, color=BLUE, size=14)

        elif n == 21:
            data = [
                ("HIMATH DHANAPALA", ["Statutory corpus + benchmark lead", "Retrieval evaluation", "High Court case-law work", "Research-paper first draft"]),
                ("LAHIRU DILSHAN", ["Case-law collection + retrieval", "OCR research", "Document-extraction research"]),
                ("PRAVEEN DE SILVA", ["BM25 search work", "Templates + output flow", "Lawyer-facing review interface"]),
            ]
            for idx, (name, bullets) in enumerate(data):
                x = 0.62+idx*4.10
                rect(slide, x, 1.88, 3.82, 4.38, fill=WHITE, line=GREY2)
                rect(slide, x, 1.88, 3.82, 0.80, fill=BLUE)
                txt(slide, name, x+0.18, 2.07, 3.45, 0.40, size=16.2, color=WHITE, bold=True)
                for j, bullet in enumerate(bullets):
                    yy = 2.95+j*0.77
                    rect(slide, x+0.19, yy+0.12, 0.08, 0.08, fill=BLUE)
                    txt(slide, bullet, x+0.36, yy, 3.20, 0.49, size=15.0)
            chip(slide, "SHARED: INTEGRATION, TESTING, AND FINAL PROJECT REVIEW", 0.62, 6.52, 12.03, 0.40,
                 fill=BLUE_PALE, color=BLUE, size=13.2)

        elif n == 22:
            logo(slide, dark=True, large=True)
            txt(slide, "Measured retrieval.\nVersioned evidence.\nLawyer-owned decisions.",
                0.72, 2.00, 11.65, 2.35, size=39, color=WHITE)
            line(slide, 0.72, 5.12, 12.26, 5.12, color=WHITE, width=0.7)
            txt(slide, "NEXT VALIDATION", 0.72, 5.39, 3.2, 0.25, size=10.5, color=WHITE, bold=True)
            txt(slide, "Attorney labels · document fields · form review · complete browser journeys",
                0.72, 5.86, 11.50, 0.72, size=20, color=WHITE)
            txt(slide, "THANK YOU", 0.72, 6.86, 4.2, 0.25, size=11, color=WHITE, bold=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)
    print(f"slides={len(prs.slides)}")


if __name__ == "__main__":
    build()
