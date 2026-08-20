"""Draftly mid-evaluation deck - layout kit matching the beige/brown template."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

TEMPLATE = r"..\White Beige and Brown Simple Minimalist Clean Final Project Presentation.pptx"

# ---- palette (lifted from the template) ----
INK    = RGBColor(0x33, 0x1B, 0x0D)   # dark brown text
BRONZE = RGBColor(0x8C, 0x77, 0x53)   # muted bronze accent
CREAM  = RGBColor(0xFA, 0xF7, 0xF0)   # page background
CARD   = RGBColor(0xF1, 0xEA, 0xDC)   # tinted card
RULE   = RGBColor(0xD9, 0xCD, 0xB8)   # hairline / border
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

REG  = "Open Sauce"
BOLD = "Open Sauce Bold"
ITAL = "Open Sauce Italics"

# ---- geometry ----
SW, SH = 20.0, 11.25
ML, MR = 1.12, 18.88          # content margins
CW = MR - ML                  # 17.76
EYEBROW_Y = 1.16
TITLE_Y   = 2.18
PAGE_Y    = 9.83
BODY_BOTTOM = 9.60            # keep content above the page number

WARNINGS = []

# ---------------------------------------------------------------- primitives
def _set_bg(slide, rgb=CREAM):
    xml = (
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<p:bgPr><a:solidFill><a:srgbClr val="{rgb}"/></a:solidFill>'
        '<a:effectLst/></p:bgPr></p:bg>'
    )
    slide.background  # touch
    cSld = slide._element.find(qn('p:cSld'))
    cSld.insert(0, etree.fromstring(xml))


def new_slide(prs, bg=CREAM):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for shp in list(s.shapes):          # blank layout still carries date/footer phs
        shp._element.getparent().remove(shp._element)
    _set_bg(s, bg)
    return s


# rough width of a string in inches, for overflow checking
_WIDE = set("MWmw@%&")
_NARROW = set("iljt.,;:'!|()[] ")
def _text_width(t, pt, bold=False):
    em = pt / 72.0
    w = 0.0
    for ch in t:
        if ch in _WIDE:      w += 0.85 * em
        elif ch in _NARROW:  w += 0.30 * em
        elif ch.isupper():   w += 0.63 * em
        elif ch.isdigit():   w += 0.55 * em
        else:                w += 0.51 * em
    return w * (1.045 if bold else 1.0)


def _wrapped_lines(text, box_w, pt, bold=False):
    words, lines, cur = text.split(), 0, ""
    for wd in words:
        trial = wd if not cur else cur + " " + wd
        if _text_width(trial, pt, bold) <= box_w:
            cur = trial
        else:
            lines += 1
            cur = wd
    return lines + (1 if cur else 0)


def textbox(slide, x, y, w, text, *, size=21, color=INK, font=REG, bold=False,
            italic=False, align=PP_ALIGN.LEFT, line=1.40, space=0.0,
            bullet=False, tag="", h=None, anchor=MSO_ANCHOR.TOP, caps_spacing=None):
    """text may be a str or a list of paragraphs. Returns (shape, height_used)."""
    paras = [text] if isinstance(text, str) else list(text)
    line_h = size * line / 72.0
    gap = space / 72.0
    n_lines = 0
    avail = w - (0.34 if bullet else 0.0)
    for p in paras:
        n_lines += _wrapped_lines(p, avail, size, bold)
    est_h = n_lines * line_h + max(0, len(paras) - 1) * gap + 0.06
    box_h = h if h is not None else est_h

    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(box_h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ptxt in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        if i:
            para.space_before = Pt(space)
        r = para.add_run()
        r.text = ("\u2022   " + ptxt) if bullet else ptxt
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
        if bullet:
            _hanging(para, 0.34)
        if caps_spacing:
            r._r.get_or_add_rPr().set('spc', str(int(caps_spacing * 100)))
    if y + est_h > BODY_BOTTOM + 0.05 and tag:
        WARNINGS.append(f"{tag}: bottom {y + est_h:.2f}\" exceeds {BODY_BOTTOM}\"")
    return tb, est_h


def _hanging(para, indent_in):
    pPr = para._pPr if para._pPr is not None else para._p.get_or_add_pPr()
    emu = str(int(indent_in * 914400))
    pPr.set('marL', emu)
    pPr.set('indent', '-' + emu)


def rect(slide, x, y, w, h, fill=CARD, line=None, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1.0)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        adj = (radius if radius is not None else 0.10) / min(w, h) * 1.0
        try:
            sp.adjustments[0] = min(0.5, adj)
        except Exception:
            pass
    sp.text_frame.text = ""
    return sp


def hline(slide, x, y, w, color=RULE, weight=1.0):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln


def vline(slide, x, y, h, color=RULE, weight=1.0):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln


def dot(slide, x, y, d=0.16, color=BRONZE):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp


# ---------------------------------------------------------------- chrome
def chrome(slide, section, page, right="Draftly | Team A"):
    textbox(slide, ML, EYEBROW_Y, 9.0, section.upper(), size=18, font=BOLD,
            bold=True, color=BRONZE, line=1.05, caps_spacing=0.6)
    textbox(slide, MR - 5.0, EYEBROW_Y, 5.0, right, size=18, font=BOLD,
            bold=True, color=BRONZE, align=PP_ALIGN.RIGHT, line=1.05)
    textbox(slide, 13.89, PAGE_Y, 4.99, str(page), size=16, color=INK,
            align=PP_ALIGN.RIGHT, line=1.05)


def title(slide, text, y=TITLE_Y, size=48, w=None, color=INK):
    return textbox(slide, ML, y, w or 15.5, text, size=size, color=color, line=1.14)


def kicker(slide, x, y, text, w=6.0, color=BRONZE, size=24, font=REG):
    return textbox(slide, x, y, w, text, size=size, color=color, font=font, line=1.20)
