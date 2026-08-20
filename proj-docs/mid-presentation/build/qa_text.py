# -*- coding: utf-8 -*-
"""Content QA: dump deck text and flag template leftovers."""
import io
import re
import sys

from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8',
                              errors='replace')

BAD = re.compile(r'xxxx|lorem|ipsum|estelle|darcy|shodwe|drew feig|eleanor|'
                 r'donna stroupe|this (page|slide)|planner|student',
                 re.I)

prs = Presentation(r"..\Draftly - Mid Evaluation.pptx")
hits = []
for i, s in enumerate(prs.slides, 1):
    parts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            parts.append(sh.text_frame.text.strip())
    joined = ' | '.join(parts)
    print("--- %02d --- %s" % (i, joined[:200]))
    for m in BAD.finditer(joined):
        hits.append((i, m.group()))

print("\nslides:", len(prs.slides))
print("leftover hits:", hits if hits else "none")
